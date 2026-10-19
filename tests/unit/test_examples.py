"""The example workbooks (examples/*.xlsx), checked without Excel.

tests/excel/test_examples.py compares each formula's result with the expected value beside it
in real Excel. These tests check what that comparison relies on: that every formula has
an expected value beside it, calls only the documented functions or the direct functions
its workbook's own script declares, and names a script and function that exist; and that
the committed workbooks are what scripts/make_examples.py writes. They need openpyxl, and
run on any system.
"""

import ast
import importlib.util
import re
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.utils import range_boundaries  # noqa: E402
from openpyxl.worksheet.formula import ArrayFormula  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
EXAMPLES = ROOT / "examples"
WORKBOOKS = sorted(EXAMPLES.glob("*.xlsx"))

# The functions docs/architecture.md documents, read from its table.
DOCUMENTED = set(re.findall(r"^\| `(Py\w+)\(", (ROOT / "docs" / "architecture.md").read_text("utf-8"), re.M))
# Excel's own functions the examples may wrap a result in, all in Excel since 2007.
EXCEL = {"ROWS", "COLUMNS", "ISNUMBER", "SEARCH"}
# The two deliberate failures in basics.xlsx.
MISSING = {"no_such_script.py", "no_such_function"}


def load_generator():
    spec = importlib.util.spec_from_file_location("make_examples", ROOT / "scripts" / "make_examples.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def checks(path):
    """Each formula on each sheet: (sheet, cell, formula, result rows, result columns)."""
    wb = openpyxl.load_workbook(path)
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, ArrayFormula):
                    left, top, right, bottom = range_boundaries(cell.value.ref)
                    yield ws, cell, cell.value.text, bottom - top + 1, right - left + 1
                elif cell.data_type == "f":
                    yield ws, cell, cell.value, 1, 1


def declared(path):
    """The direct functions the workbook's own script marks @pyinex.func, by the names Excel
    is given, if the workbook loads that script with PyLoad; otherwise none."""
    script = path.with_suffix(".py")
    if not any(f'PyLoad("{script.name}")' in formula for _, _, formula, _, _ in checks(path)):
        return set()
    names = set()
    for node in ast.parse(script.read_text("utf-8")).body:
        if not isinstance(node, (ast.FunctionDef, ast.ClassDef)):
            continue
        for mark in node.decorator_list:
            call = mark if isinstance(mark, ast.Call) else None
            if ast.unparse(call.func if call else mark) != "pyinex.func":
                continue
            given = [k.value.value for k in call.keywords if k.arg == "name"] if call else []
            names.add(given[0] if given else node.name)
    return names


def cell_contents(path):
    wb = openpyxl.load_workbook(path)
    return {
        ws.title: {c.coordinate: (c.data_type, getattr(c.value, "text", c.value), getattr(c.value, "ref", None)) for row in ws.iter_rows() for c in row}
        for ws in wb
    }


def test_the_documented_functions_are_found():
    assert DOCUMENTED == {"PyCall", "PyConsole", "PyReloadOnChange", "PyInfo", "PyLoad"}


def test_there_is_a_workbook_for_each_script():
    assert [p.stem for p in WORKBOOKS] == sorted(p.stem for p in EXAMPLES.glob("*.py"))
    assert WORKBOOKS


@pytest.mark.parametrize("path", WORKBOOKS, ids=lambda p: p.name)
def test_every_formula_has_an_expected_value_beside_it(path):
    found = list(checks(path))
    assert found, "no formulas"
    for ws, cell, formula, rows, columns in found:
        expected = [
            ws.cell(cell.row + r, cell.column + columns + c) for r in range(rows) for c in range(columns)
        ]
        assert any(e.value is not None for e in expected), f"{ws.title}!{cell.coordinate}: no expected value"
        for e in expected:
            assert e.data_type != "f" and not isinstance(e.value, ArrayFormula), f"{ws.title}!{e.coordinate}"


@pytest.mark.parametrize("path", WORKBOOKS, ids=lambda p: p.name)
def test_every_formula_calls_only_documented_functions(path):
    direct = declared(path)
    for ws, cell, formula, _, _ in checks(path):
        names = set(re.findall(r"([A-Za-z_][\w.]*)\(", re.sub(r'"[^"]*"', '""', formula)))
        assert names & (DOCUMENTED | direct), f"{ws.title}!{cell.coordinate}: {formula} calls no Pyinex function"
        assert names <= DOCUMENTED | EXCEL | direct, f"{ws.title}!{cell.coordinate}: {formula}"


def test_a_direct_function_is_accepted_only_from_the_workbooks_own_loaded_script():
    assert declared(EXAMPLES / "direct_functions.xlsx") == {
        "Discount", "CompoundGrowth", "Total", "Where", "UsedRows", "FilledCells"
    }
    assert declared(EXAMPLES / "ranges.xlsx") == set()  # It loads no script with PyLoad.


@pytest.mark.parametrize("path", WORKBOOKS, ids=lambda p: p.name)
def test_every_pycall_names_a_script_and_function_that_exist(path):
    for ws, cell, formula, _, _ in checks(path):
        for script, function in re.findall(r'PyCall\("([^"]+)","([^"]+)"', formula):
            if script in MISSING or function in MISSING:
                continue
            assert script == path.with_suffix(".py").name, f"{ws.title}!{cell.coordinate}: {formula}"
            tree = ast.parse((EXAMPLES / script).read_text("utf-8"))
            defined = {n.name for n in tree.body if isinstance(n, ast.FunctionDef)} | {
                t.id for n in tree.body if isinstance(n, ast.Assign) for t in n.targets if isinstance(t, ast.Name)
            }
            assert function in defined, f"{ws.title}!{cell.coordinate}: {function} is not in {script}"
        for script in re.findall(r'PyLoad\("([^"]+)"', formula):
            assert script == path.with_suffix(".py").name, f"{ws.title}!{cell.coordinate}: {formula}"


def test_the_workbooks_are_what_the_generator_writes(tmp_path):
    for fresh in load_generator().build(tmp_path):
        committed = EXAMPLES / fresh.name
        assert committed.is_file(), f"run scripts/make_examples.py: {committed.name} is missing"
        assert cell_contents(committed) == cell_contents(fresh), f"run scripts/make_examples.py: {fresh.name} is stale"
